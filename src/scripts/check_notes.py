"""
scripts/check_notes.py

Diagnostic: directly inspects has_notes_slide + notes text for every
slide, bypassing the extraction pipeline entirely. Run this against
your actual .pptx to settle whether missing "notes" in the JSON means
"no notes exist" or "something's being missed."

Usage:
    python scripts/check_notes.py path/to/deck.pptx
"""

import sys
from pptx import Presentation


def main(pptx_path: str):
    prs = Presentation(pptx_path)
    for i, slide in enumerate(prs.slides, start=1):
        has_notes = slide.has_notes_slide
        text = ""
        if has_notes:
            text = slide.notes_slide.notes_text_frame.text.strip()
        print(f"Slide {i}: has_notes_slide={has_notes} | text={text!r}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python scripts/check_notes.py <path-to-deck.pptx>")
        raise SystemExit(1)
    main(sys.argv[1])
